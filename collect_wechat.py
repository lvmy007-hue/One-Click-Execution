# -*- coding: utf-8 -*-
"""无人值守：按到达本机时间收集店表。每天 09:30 截止，不看文件名日期。

09:30 前到达 → 当天目录；09:30 及以后 → 次日目录。
"""
from __future__ import annotations

import argparse
import atexit
import hashlib
import json
import os
import shutil
import subprocess
import sys
import time
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).resolve().parent
DEFAULT_CFG = HERE / "config.json"
LOG_FILE = HERE / "collect.log"
DAEMON_PID_FILE = HERE / "collect_daemon.pid"
SUP_PID_FILE = HERE / "collect_supervise.pid"
CREATE_NO_WINDOW = 0x08000000
EXCEL_EXT = {".xlsx", ".xls"}
NAME_BLOCK = (
    "查邮箱", "查地址", "好评", "差评", "每日数据", "内部cyx", "链接查",
    "精准", "联系方式", "匹配成功", "匹配失败", "精简版",
)
FOUND_SKIP = ("链接查", "链接", "订单查", "好评", "差评", "每日数据", "内部cyx", "精准", "联系方式")
HEADER_MARKS = ("店铺", "买家全名", "城市", "州", "邮编")
DEFAULT_MATCH_ROOT = r"D:\桌面\地址匹配源文件\匹配数据"
FOUND_SUBDIR = "收表汇总"


def log(*args):
    line = " ".join(str(a) for a in args)
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    text = f"{stamp} {line}"
    try:
        sys.stdout.write(text + "\n")
        sys.stdout.flush()
    except Exception:
        pass
    try:
        with LOG_FILE.open("a", encoding="utf-8") as f:
            f.write(text + "\n")
    except OSError:
        pass


def load_cfg(path: Path) -> dict:
    cfg = {
        "wechat_root": r"D:\文档\xwechat_files",
        "drop_dir": r"D:\桌面\地址匹配源文件\待收集",
        "dest_root": r"D:\桌面\地址匹配源文件\原始数据",
        "poll_seconds": 2,
        "stable_seconds": 1.2,
        "collect_cutoff": "09:30",
        "found_cutoff": "17:40",
        "match_dest_root": DEFAULT_MATCH_ROOT,
        "need_headers": list(HEADER_MARKS),
    }
    if path.is_file():
        cfg.update(json.loads(path.read_text(encoding="utf-8")))
    return cfg


def day_start(dt: datetime | None = None) -> datetime:
    dt = dt or datetime.now()
    return datetime(dt.year, dt.month, dt.day)


def month_dir_name(dt: datetime) -> str:
    return f"{dt.month}月"


def date_dir_name(dt: datetime) -> str:
    return f"{dt.month}.{dt.day:02d}"


def parse_hm(value, default_h: int = 9, default_m: int = 30) -> tuple[int, int]:
    if value is None or value == "":
        return default_h, default_m
    if isinstance(value, bool):
        return default_h, default_m
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        h = int(value)
        m = int(round((float(value) - h) * 60)) if isinstance(value, float) else 0
        return max(0, min(23, h)), max(0, min(59, m))
    s = str(value).strip().replace("：", ":")
    try:
        if ":" in s:
            a, b = s.split(":", 1)
            return max(0, min(23, int(a))), max(0, min(59, int(b)))
        return max(0, min(23, int(s))), 0
    except (TypeError, ValueError):
        return default_h, default_m


def fmt_hm(hour: int, minute: int = 0) -> str:
    return f"{hour:02d}:{minute:02d}"


def cutoff_hm(cfg: dict) -> tuple[int, int]:
    return parse_hm(cfg.get("collect_cutoff") or cfg.get("collect_cutoff_hour"), 9, 30)


def found_cutoff_hm(cfg: dict) -> tuple[int, int]:
    return parse_hm(cfg.get("found_cutoff") or cfg.get("found_cutoff_hour"), 17, 40)


def cutoff_hour(cfg: dict) -> int:
    return cutoff_hm(cfg)[0]


def found_cutoff_hour(cfg: dict) -> int:
    return found_cutoff_hm(cfg)[0]


def batch_day(arrived: datetime, hour: int, minute: int = 0) -> datetime:
    """到达时刻归属的业务日：截止时刻之前算当天，到点及之后算次日。"""
    d = day_start(arrived)
    cut = d.replace(hour=hour, minute=minute, second=0, microsecond=0)
    if arrived >= cut:
        return d + timedelta(days=1)
    return d


def collect_day(cfg: dict, now: datetime | None = None) -> datetime:
    """当前正在往哪个业务日目录里收新表。"""
    now = now or datetime.now()
    h, m = cutoff_hm(cfg)
    return batch_day(now, h, m)


def window_since(day: datetime, hour: int, minute: int = 0) -> datetime:
    prev = day - timedelta(days=1)
    return datetime(prev.year, prev.month, prev.day, hour, minute, 0)


def dest_dir(cfg: dict, dt: datetime | None = None) -> Path:
    dt = dt or datetime.now()
    d = Path(cfg["dest_root"]) / month_dir_name(dt) / date_dir_name(dt)
    d.mkdir(parents=True, exist_ok=True)
    return d


def collect_dest(cfg: dict, now: datetime | None = None) -> Path:
    return dest_dir(cfg, collect_day(cfg, now))


def found_collect_day(cfg: dict, now: datetime | None = None) -> datetime:
    """查出表：截止时刻前归当天匹配目录，到点及以后归次日。"""
    return found_dest_day(cfg, now or datetime.now())


def found_rule_start(cfg: dict) -> datetime:
    raw = str(cfg.get("found_rule_start") or "2026-10-01")[:10]
    try:
        return datetime.strptime(raw, "%Y-%m-%d")
    except ValueError:
        return datetime(2026, 10, 1)


def found_dest_day(cfg: dict, arrived: datetime) -> datetime:
    """规则生效日前到达的查出表，全部归生效日目录（10.01 起才按查出截止切日）。"""
    start = found_rule_start(cfg)
    if arrived < start:
        return start
    h, m = found_cutoff_hm(cfg)
    return batch_day(arrived, h, m)


def match_dest_dir(cfg: dict, dt: datetime | None = None) -> Path:
    dt = dt or datetime.now()
    root = Path(cfg.get("match_dest_root") or DEFAULT_MATCH_ROOT)
    d = root / month_dir_name(dt) / date_dir_name(dt)
    d.mkdir(parents=True, exist_ok=True)
    return d


def found_zip_path(dest: Path, dt: datetime) -> Path:
    return dest / f"{date_dir_name(dt)}数据.zip"


def found_tables_dir(dest: Path) -> Path:
    d = dest / FOUND_SUBDIR
    d.mkdir(parents=True, exist_ok=True)
    return d


def is_found_name(name: str) -> bool:
    if any(b in name for b in FOUND_SKIP):
        return False
    if "查出" not in name:
        return False
    return "查地址" in name or "查邮箱" in name


def is_found_table(path: Path) -> bool:
    return path.suffix.lower() in EXCEL_EXT and is_found_name(path.name)


def gather_found_xlsx(dest: Path) -> list[Path]:
    """查出表进「收表汇总」。根目录残留的也迁进去。"""
    folder = found_tables_dir(dest)
    for p in list(dest.glob("*.xlsx")):
        if not is_found_name(p.name):
            continue
        target = folder / p.name
        if target.exists():
            if target.stat().st_size == p.stat().st_size:
                try:
                    p.unlink()
                except OSError:
                    pass
                continue
            target = unique_dest_path(folder, p.name)
        try:
            p.replace(target)
        except OSError:
            shutil.copy2(p, target)
            try:
                p.unlink()
            except OSError:
                pass
    return sorted(p for p in folder.glob("*.xlsx") if is_found_name(p.name))


def rebuild_found_zip(dest: Path, dt: datetime) -> Path | None:
    files = gather_found_xlsx(dest)
    zpath = found_zip_path(dest, dt)
    if not files:
        if zpath.exists():
            zpath.unlink()
        return None
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in files:
            zf.write(p, p.name)
    return zpath


def wechat_file_dirs(root: str) -> list[Path]:
    base = Path(root)
    if not base.is_dir():
        return []
    return [p / "msg" / "file" for p in base.glob("wxid_*") if (p / "msg" / "file").is_dir()]


def month_stamp_dirs(file_root: Path, now: datetime) -> list[Path]:
    y, m = now.year, now.month
    found = []
    for _ in range(2):
        d = file_root / f"{y:04d}-{m:02d}"
        if d.is_dir():
            found.append(d)
        m -= 1
        if m <= 0:
            m, y = 12, y - 1
    return found or [file_root]


def iter_excel(folder: Path):
    if not folder.is_dir():
        return
    for p in folder.iterdir():
        if p.is_file() and p.suffix.lower() in EXCEL_EXT:
            yield p


def name_blocked(name: str) -> bool:
    low = name.lower()
    return any(k in low for k in NAME_BLOCK)


def read_header(path: Path) -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(max_row=5, values_only=True):
            vals = [str(c).strip() for c in row if c is not None and str(c).strip() and str(c) != "None"]
            if any(m in vals for m in HEADER_MARKS):
                wb.close()
                return vals
        wb.close()
    except Exception:
        return []
    return []


def is_store_table(path: Path, need_headers: list[str]) -> tuple[bool, str]:
    if path.suffix.lower() not in EXCEL_EXT:
        return False, "不是表格"
    if name_blocked(path.name):
        return False, "文件名排除"
    header = read_header(path)
    if not header:
        return False, "读不到表头"
    hit = [h for h in need_headers if h in header]
    if len(hit) < 4:
        return False, "表头不像店表"
    return True, "ok"


def arrived_at(path: Path) -> datetime:
    """本机首次出现时间（创建时间）。不看文件名里的 9.21 / 20260924。"""
    st = path.stat()
    return datetime.fromtimestamp(st.st_ctime)


def wait_stable(path: Path, seconds: float) -> bool:
    try:
        s1 = path.stat().st_size
    except OSError:
        return False
    time.sleep(max(0.05, seconds))
    try:
        st = path.stat()
    except OSError:
        return False
    return st.st_size == s1 and st.st_size > 0


def sha1_file(path: Path) -> str:
    h = hashlib.sha1()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def dest_has_same(dest: Path, src: Path) -> Path | None:
    same_name = dest / src.name
    if same_name.is_file() and same_name.stat().st_size == src.stat().st_size:
        return same_name
    src_size = src.stat().st_size
    for p in dest.glob("*.xlsx"):
        if p.stat().st_size != src_size:
            continue
        try:
            if sha1_file(p) == sha1_file(src):
                return p
        except OSError:
            continue
    return None


def unique_dest_path(dest: Path, name: str) -> Path:
    target = dest / name
    if not target.exists():
        return target
    stem, ext = os.path.splitext(name)
    n = 1
    while True:
        cand = dest / f"{stem}({n}){ext}"
        if not cand.exists():
            return cand
        n += 1


def copy_one(src: Path, dest: Path, dry: bool) -> tuple[str, str]:
    existed = dest_has_same(dest, src)
    if existed:
        return "skip", f"已有 {existed.name}"
    target = unique_dest_path(dest, src.name)
    if dry:
        return "dry", f"将复制 → {target.name}"
    shutil.copy2(src, target)
    return "copy", src.name if target.name == src.name else f"{src.name} → {target.name}"


def collect_candidates(cfg: dict, since: datetime, extra: list[Path]) -> list[Path]:
    seen: set[str] = set()
    out: list[Path] = []

    def add(p: Path, require_since: bool):
        key = str(p.resolve()).lower()
        if key in seen:
            return
        if require_since and arrived_at(p) < since:
            return
        seen.add(key)
        out.append(p)

    now = datetime.now()
    for file_root in wechat_file_dirs(cfg["wechat_root"]):
        for month_dir in month_stamp_dirs(file_root, now):
            for p in iter_excel(month_dir):
                add(p, True)

    drop = Path(cfg["drop_dir"])
    drop.mkdir(parents=True, exist_ok=True)
    for p in iter_excel(drop):
        add(p, True)

    for folder in extra:
        for p in iter_excel(folder):
            add(p, False)
    return out


def harvest(cfg: dict, since: datetime, extra: list[Path], dry: bool, quiet_skip: bool) -> dict:
    h, m = cutoff_hm(cfg)
    dest_day = collect_day(cfg)
    dest = dest_dir(cfg, dest_day)
    stats = {"copy": 0, "skip": 0, "reject": 0, "dry": 0}
    need = cfg.get("need_headers") or list(HEADER_MARKS)
    wait = 0.08 if dry else 0.25
    for src in collect_candidates(cfg, since, extra):
        if batch_day(arrived_at(src), h, m) != dest_day:
            continue
        if not wait_stable(src, wait):
            continue
        ok, why = is_store_table(src, need)
        if not ok:
            stats["reject"] += 1
            continue
        action, msg = copy_one(src, dest, dry)
        stats[action] = stats.get(action, 0) + 1
        if action == "skip" and quiet_skip:
            continue
        log(("DRY " if dry else "") + msg)
    return stats


def harvest_found(cfg: dict, extra: list[Path], dry: bool, quiet_skip: bool) -> dict:
    h, m = found_cutoff_hm(cfg)
    now_day = found_collect_day(cfg)
    since = window_since(now_day - timedelta(days=1), h, m)
    stats = {"copy": 0, "skip": 0, "reject": 0, "dry": 0}
    wait = 0.08 if dry else 0.25
    touched: set[datetime] = set()
    for src in collect_candidates(cfg, since, extra):
        if not is_found_table(src):
            continue
        if not wait_stable(src, wait):
            continue
        day = found_dest_day(cfg, arrived_at(src))
        dest = found_tables_dir(match_dest_dir(cfg, day))
        action, msg = copy_one(src, dest, dry)
        stats[action] = stats.get(action, 0) + 1
        if action == "copy":
            touched.add(day)
        if action == "skip" and quiet_skip:
            continue
        log(("DRY " if dry else "") + "查出 " + msg + " → " + date_dir_name(day))
    if not dry:
        for day in {now_day, now_day - timedelta(days=1)} | touched:
            dest = match_dest_dir(cfg, day)
            zp = rebuild_found_zip(dest, day)
            if zp:
                log("查出 ZIP", zp, zp.stat().st_size, "字节")
    return stats


def run_once(cfg: dict, extra: list[Path], dry: bool) -> dict:
    h, m = cutoff_hm(cfg)
    fh, fm = found_cutoff_hm(cfg)
    dest_day = collect_day(cfg)
    dest = dest_dir(cfg, dest_day)
    since = window_since(dest_day, h, m)
    log(f"规则: 店表每天 {fmt_hm(h, m)} 截止，不看文件名日期")
    log(f"本窗口 {since.strftime('%m-%d %H:%M')} ～ {dest_day.strftime('%m-%d')} {fmt_hm(h, m)} → {date_dir_name(dest_day)}")
    log(f"微信库: {cfg['wechat_root']}")
    for d in wechat_file_dirs(cfg["wechat_root"]):
        log(f"  {d}")
    log(f"汇总到: {dest}")
    stats = harvest(cfg, since, extra, dry, quiet_skip=False)
    found_stats = harvest_found(cfg, extra, dry, quiet_skip=False)
    files_now = sorted(p.name for p in dest.glob("*.xlsx"))
    log(f"复制 {stats['copy']}  跳过 {stats['skip']}  排除 {stats['reject']}")
    log(f"当天目录 {len(files_now)} 个")
    for name in files_now:
        log(f"  {name}")
    log(f"查出复制 {found_stats['copy']}  跳过 {found_stats['skip']}  截止 {fmt_hm(fh, fm)}")
    return {"dest": str(dest), "files": files_now, **stats}


def watch(cfg: dict, extra: list[Path], daemon: bool):
    write_pid(DAEMON_PID_FILE)
    atexit.register(lambda: clear_pid(DAEMON_PID_FILE, os.getpid()))
    need_done: set[str] = set()
    poll = float(cfg.get("poll_seconds", 2))
    h, m = cutoff_hm(cfg)
    fh, fm = found_cutoff_hm(cfg)
    last_batch = collect_day(cfg)
    last_found = found_collect_day(cfg)
    dest = dest_dir(cfg, last_batch)
    log("无人值守监听已启动")
    log(f"每天 {fmt_hm(h, m)} 截止，当前归档: {dest}")
    log(f"查出表每天 {fmt_hm(fh, fm)} 截止，当前归档: {match_dest_dir(cfg, last_found)}")
    log("表格要落到硬盘，需微信保持登录，并开启「文件自动下载」（店表很小，会自动下）")
    harvest(cfg, window_since(last_batch, h, m), extra, False, quiet_skip=True)
    harvest_found(cfg, extra, False, quiet_skip=True)
    try:
        while True:
            h, m = cutoff_hm(cfg)
            fh, fm = found_cutoff_hm(cfg)
            batch = collect_day(cfg)
            if batch != last_batch:
                last_batch = batch
                need_done.clear()
                dest = dest_dir(cfg, batch)
                log(f"已过 {fmt_hm(h, m)} 截止，改归档到 {dest}")
            found_batch = found_collect_day(cfg)
            if found_batch != last_found:
                last_found = found_batch
                log(f"已过 {fmt_hm(fh, fm)} 截止，查出改归档到 {match_dest_dir(cfg, found_batch)}")
            dest = dest_dir(cfg, batch)
            since = window_since(batch, h, m)
            found_since = window_since(found_collect_day(cfg) - timedelta(days=1), fh, fm)
            scan_since = min(since, found_since)
            try:
                for src in collect_candidates(cfg, scan_since, extra):
                    key = str(src.resolve()).lower()
                    if key in need_done:
                        continue
                    if not wait_stable(src, cfg.get("stable_seconds", 1.2)):
                        continue
                    if is_found_table(src):
                        need_done.add(key)
                        fday = found_dest_day(cfg, arrived_at(src))
                        fdest = match_dest_dir(cfg, fday)
                        action, msg = copy_one(src, found_tables_dir(fdest), False)
                        if action != "skip":
                            log("查出 " + msg + " → " + date_dir_name(fday))
                            zp = rebuild_found_zip(fdest, fday)
                            if zp:
                                log("查出 ZIP", zp.name, zp.stat().st_size, "字节")
                        continue
                    if batch_day(arrived_at(src), h, m) != batch:
                        continue
                    ok, why = is_store_table(src, cfg.get("need_headers") or list(HEADER_MARKS))
                    need_done.add(key)
                    if not ok:
                        continue
                    action, msg = copy_one(src, dest, False)
                    if action != "skip":
                        log(msg)
            except Exception as e:
                log("本轮收集出错，继续监听:", e)
            time.sleep(poll)
    except KeyboardInterrupt:
        log("已停止监听")


def pid_alive(pid: int) -> bool:
    if pid <= 0:
        return False
    if os.name == "nt":
        import ctypes
        PROCESS_QUERY_LIMITED_INFORMATION = 0x1000
        handle = ctypes.windll.kernel32.OpenProcess(PROCESS_QUERY_LIMITED_INFORMATION, False, int(pid))
        if handle:
            ctypes.windll.kernel32.CloseHandle(handle)
            return True
        return False
    try:
        os.kill(pid, 0)
    except PermissionError:
        return True
    except OSError:
        return False
    return True


def read_pid(path: Path) -> int:
    try:
        return int(path.read_text(encoding="utf-8").strip())
    except Exception:
        return 0


def write_pid(path: Path, pid: int | None = None) -> None:
    path.write_text(str(pid or os.getpid()), encoding="utf-8")


def clear_pid(path: Path, pid: int | None = None) -> None:
    cur = read_pid(path)
    if pid is not None and cur and cur != pid:
        return
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def has_other(role: str) -> bool:
    path = SUP_PID_FILE if role == "supervise" else DAEMON_PID_FILE
    pid = read_pid(path)
    return bool(pid) and pid != os.getpid() and pid_alive(pid)


def start_daemon() -> subprocess.Popen:
    script = HERE / "collect_wechat.py"
    return subprocess.Popen(
        [str(pythonw_path()), str(script), "--daemon"],
        cwd=str(HERE),
        close_fds=True,
        creationflags=CREATE_NO_WINDOW,
    )


def supervise(cfg: dict) -> None:
    if has_other("supervise"):
        log("守护进程已在运行，本次退出")
        return
    write_pid(SUP_PID_FILE)
    atexit.register(lambda: clear_pid(SUP_PID_FILE, os.getpid()))
    log("守护已启动：收集器崩溃后会自动拉起")
    try:
        while True:
            if has_other("daemon"):
                time.sleep(3)
                continue
            log("正在拉起收集器")
            proc = start_daemon()
            write_pid(DAEMON_PID_FILE, proc.pid)
            code = proc.wait()
            log("收集器退出", code, "，即将重新拉起")
            clear_pid(DAEMON_PID_FILE, proc.pid)
            time.sleep(2)
    except KeyboardInterrupt:
        log("守护已停止")


def pythonw_path() -> Path:
    exe = Path(sys.executable)
    alt = exe.with_name("pythonw.exe")
    return alt if alt.is_file() else exe


def install_startup() -> Path:
    startup = Path(os.environ["APPDATA"]) / r"Microsoft\Windows\Start Menu\Programs\Startup"
    startup.mkdir(parents=True, exist_ok=True)
    vbs = startup / "微信店表无人值守.vbs"
    pyw = pythonw_path()
    script = HERE / "collect_wechat.py"
    line = f'"{pyw}" "{script}" --supervise'
    vbs.write_text(
        "Set s = CreateObject(\"WScript.Shell\")\r\n"
        f"s.CurrentDirectory = \"{HERE}\"\r\n"
        f"s.Run \"\"\"{pyw}\"\" \"\"{script}\"\" --supervise\", 0, False\r\n",
        encoding="utf-8",
    )
    log(f"已写入开机启动: {vbs}")
    return vbs


def parse_args():
    ap = argparse.ArgumentParser(description="按当日到达收集微信店表")
    ap.add_argument("--config", default=str(DEFAULT_CFG))
    ap.add_argument("--once", action="store_true")
    ap.add_argument("--watch", action="store_true")
    ap.add_argument("--daemon", action="store_true", help="无人值守：写日志、跨日切换")
    ap.add_argument("--supervise", action="store_true", help="守护：收集器崩溃后自动拉起")
    ap.add_argument("--install-startup", action="store_true", help="安装开机自启")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--extra", action="append", default=[])
    return ap.parse_args()


def main():
    args = parse_args()
    cfg = load_cfg(Path(args.config))
    extra = [Path(p) for p in args.extra]
    if args.install_startup:
        install_startup()
        return
    if args.once or args.dry_run:
        run_once(cfg, extra, args.dry_run)
        return
    if args.supervise:
        supervise(cfg)
        return
    if has_other("daemon"):
        log("收集器已在运行，本次不重复启动")
        return
    watch(cfg, extra, daemon=args.daemon or args.watch or True)


if __name__ == "__main__":
    main()
